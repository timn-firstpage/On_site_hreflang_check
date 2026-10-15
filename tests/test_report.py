import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from build_report import build_report, report_rows, validate
from openpyxl import load_workbook


def sample():
    return json.loads((Path(__file__).resolve().parents[1] / 'examples/findings.json').read_text(encoding='utf-8'))


class ReportTests(unittest.TestCase):
    def test_large_missing_review_has_no_confirmed_issue_rows(self):
        data = sample()
        for c in data['checks']:
            c.update(result='√', evidence_state='available', issues=[])
        data['checks'][2].update(result='Human check', finding='发现 101 个缺少 hreflang 的 URL，需复核。', instruction='请检查是否有大量参数页或 Screaming Frog 抓取异常。')
        rows = report_rows(data)
        self.assertEqual(list(rows), ['Checklist'])
        self.assertEqual(rows['Checklist'][3][2], 'Human check')
        self.assertIn('101', rows['Checklist'][3][3])
        self.assertEqual(rows['Checklist'][3][4], data['checks'][2]['instruction'])

    def test_confirmed_other_filter_issue_can_coexist_with_missing_review(self):
        data = sample()
        issue = {'address': 'https://example.com/en/', 'issue': '缺少返回链接；目标：https://example.com/zh/', 'instruction': '请添加指向对应中文页面的 hreflang 链接。'}
        data['checks'][2].update(result='X', evidence_state='available', finding='发现缺少返回链接；另有 101 个 Missing URL 待复核。', instruction='请补充返回链接，并检查参数页或 SF 抓取异常。', issues=[issue])
        rows = report_rows(data)
        self.assertEqual(rows['Checklist'][3][2], 'X')
        # One 15.2 issue and the confirmed 15.3 return-link defect; no 101 raw Missing rows.
        self.assertEqual(len(rows['15. HREFLANG']), 3)
        self.assertIn('缺少返回链接', rows['15. HREFLANG'][2][1])

    def test_checklist_separates_instruction_without_coverage(self):
        data = sample()
        for c in data['checks']:
            c.pop('coverage', None)
            c['instruction'] = '请补充 SF 导出。' if c['id'] == '15.3' else '无需处理。'
        rows = report_rows(data)['Checklist']
        self.assertEqual(rows[0], ['Item No.', 'Item Name', 'Check Result', 'Findings', 'Instruction'])
        self.assertEqual(rows[3][4], '请补充 SF 导出。')
        self.assertNotIn('请补充 SF 导出。', rows[3][3])

    def test_checklist_requires_human_next_step(self):
        data = sample()
        for c in data['checks']:
            c['instruction'] = '无需处理。'
        data['checks'][1]['instruction'] = ''
        with self.assertRaises(ValueError):
            validate(data)

    def test_all_pass_has_no_issue_sheet(self):
        data = sample()
        for c in data['checks']:
            c.update(result='√', evidence_state='available', issues=[])
        self.assertEqual(list(report_rows(data)), ['Checklist'])

    def test_all_x_have_detail(self):
        data = sample()
        issue = data['checks'][1]['issues'][0]
        for c in data['checks']:
            c.update(result='X', evidence_state='available', issues=[copy.deepcopy(issue)])
        self.assertEqual(len(report_rows(data)['15. HREFLANG']), 4)

    def test_missing_results_are_disclosed(self):
        data = sample()
        row = report_rows(data)['Checklist'][3]
        self.assertIn('尚未取得', row[3])
        self.assertIn('未完成验证', row[3])
        self.assertEqual(row[4], '请补充六项 SF 导出。')
        del data['checks'][2]['missing_evidence_note']
        with self.assertRaises(ValueError):
            validate(data)

    def test_invalid_status_issue_and_id(self):
        for mutate in [lambda d: d['checks'][1].update(issues=[]),
                       lambda d: d['checks'][0].update(id='15.2'),
                       lambda d: d['checks'][0].update(result='Y'),
                       lambda d: d['checks'][0].update(result='Yes'),
                       lambda d: d['checks'][0].update(result='No'),
                       lambda d: d['checks'][0].update(evidence_state='unavailable')]:
            data = sample()
            mutate(data)
            with self.assertRaises(ValueError):
                validate(data)

    def test_single_language_and_explicit_skip(self):
        data = sample()
        for c in data['checks']:
            c.update(result='N/A', evidence_state='not_applicable', issues=[], finding='确认单语言；15.3 用户明确跳过。')
        self.assertEqual(list(report_rows(data)), ['Checklist'])

    def test_literal_text_roundtrip_and_no_overwrite(self):
        data = sample()
        data['checks'][0]['finding'] = '=HYPERLINK("bad","中文")'
        run_root = Path(__file__).resolve().parents[1] / 'runs'
        run_root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=run_root) as directory:
            path = build_report(data, directory)
            self.assertEqual(path.name, '15.1-15.3_Demo Company_onsite_hreflang_audit_2026-10-08.xlsx')
            wb = load_workbook(path)
            self.assertEqual(wb['Checklist']['D2'].value, data['checks'][0]['finding'])
            self.assertEqual(wb['Checklist']['D2'].data_type, 's')
            self.assertEqual(wb['Checklist']['E1'].value, 'Instruction')
            self.assertEqual(wb['Checklist']['E2'].value, '无需处理。')
            self.assertEqual(wb.sheetnames, ['Checklist', '15. HREFLANG'])
            wb.close()
            with self.assertRaises(FileExistsError):
                build_report(data, directory)


if __name__ == '__main__':
    unittest.main()
