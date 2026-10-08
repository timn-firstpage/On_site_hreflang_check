"""Validate agent-reviewed findings and export the agreed report. No crawling."""
import argparse
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import re
import unicodedata

IDS = ['15.1', '15.2', '15.3']
CHECK_HEADERS = ['Item No.', 'Item Name', 'Check Result', 'Findings', 'Coverage']
ISSUE_HEADERS = ['Address', 'Issue', 'Instruction']
STATES = {'available', 'partial', 'unavailable', 'not_applicable'}


def require_text(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'{label}: nonempty text required')
    if len(value.encode('utf-16-le', errors='surrogatepass')) // 2 > 32767:
        raise ValueError(f'{label}: Excel cell limit exceeded; archive full evidence separately')
    if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', value):
        raise ValueError(f'{label}: unsupported XML character')
    return value


def validate(data):
    if not isinstance(data, dict):
        raise ValueError('findings must be an object')
    for key in ('site_name', 'site_url', 'audit_date'):
        require_text(data.get(key), key)
    if date.fromisoformat(data['audit_date']).isoformat() != data['audit_date']:
        raise ValueError('audit_date must be YYYY-MM-DD')
    checks = data.get('checks')
    if not isinstance(checks, list) or len(checks) != 3 or not all(isinstance(c, dict) for c in checks):
        raise ValueError('checks must contain three objects')
    if sorted(str(c.get('id')) for c in checks) != IDS:
        raise ValueError('15.1, 15.2 and 15.3 must appear exactly once')
    for c in checks:
        for key in ('item_name', 'finding', 'coverage', 'result', 'evidence_state'):
            require_text(c.get(key), f"{c['id']}.{key}")
        if c['result'] not in {'√', 'X', 'N/A', 'Human check'} or c['evidence_state'] not in STATES:
            raise ValueError('Unknown result or evidence_state')
        if not isinstance(c.get('evidence'), list) or not all(isinstance(x, str) for x in c['evidence']):
            raise ValueError('evidence must be a list of references')
        if c['result'] == 'N/A' and c['evidence_state'] != 'not_applicable':
            raise ValueError('N/A requires not_applicable evidence_state')
        if c['result'] != 'N/A' and c['evidence_state'] == 'not_applicable':
            raise ValueError('not_applicable requires N/A')
        if c['result'] == '√' and c['evidence_state'] != 'available':
            if c['id'] != '15.3':
                raise ValueError('Only 15.3 allows a pass with partial/unavailable evidence')
            require_text(c.get('missing_evidence_note'), 'missing_evidence_note')
        issues = c.get('issues')
        if not isinstance(issues, list) or bool(issues) != (c['result'] == 'X'):
            raise ValueError('Every X requires issues; other results must have no issues')
        for issue in issues:
            if not isinstance(issue, dict):
                raise ValueError('issue must be an object')
            for key in ('address', 'issue', 'instruction'):
                require_text(issue.get(key), key)
    by_id = {c['id']: c for c in checks}
    if by_id['15.1']['result'] == 'N/A' and by_id['15.2']['result'] != 'N/A':
        raise ValueError('Confirmed single-language 15.1 N/A requires 15.2 N/A')


def report_rows(data):
    validate(data)
    rows = {'Checklist': [CHECK_HEADERS]}
    for c in sorted(data['checks'], key=lambda c: c['id']):
        finding, coverage = c['finding'], c['coverage']
        if c['id'] == '15.3' and c['result'] == '√' and c['evidence_state'] != 'available':
            finding += '\n' + c['missing_evidence_note']
            coverage += '\n未完成完整验证 / Not fully verified (first-check policy).'
        rows['Checklist'].append([c['id'], c['item_name'], c['result'], finding, coverage])
        for issue in c['issues']:
            rows.setdefault('15. HREFLANG', [ISSUE_HEADERS]).append([
                issue['address'], c['id'] + ' — ' + issue['issue'], issue['instruction']])
    return rows


def build_report(data, output_dir):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rows = report_rows(data)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', data['site_name']).strip(' .')
    filename = f"Onsite_hreflang_{name}_{data['audit_date']}.xlsx"
    if not name or len(filename.encode('utf-8')) > 240:
        raise ValueError('Provide a shorter usable company name')
    target = Path(output_dir) / filename
    if target.exists():
        raise FileExistsError(f'{target}: use a new run directory')
    wb = Workbook()
    wb.remove(wb.active)
    for title, table in rows.items():
        ws = wb.create_sheet(title)
        widths = [12, 47, 19, 70, 65] if title == 'Checklist' else [55, 75, 75]
        ws.freeze_panes = 'A2'
        ws.sheet_view.showGridLines = False
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for r, values in enumerate(table, 1):
            height = 32
            for col, (value, width) in enumerate(zip(values, widths), 1):
                require_text(value, f'{title}!{r},{col}')
                cell = ws.cell(r, col, value)
                cell.data_type = 's'
                cell.font = Font(name='Calibri', size=11, bold=r == 1, color='FFFFFF' if r == 1 else '223344')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.fill = PatternFill('solid', fgColor='16354A' if r == 1 else ('F0F5F8' if r % 2 == 0 else 'FFFFFF'))
                lines = sum(max(1, (sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in line) + width - 4) // (width - 3)) for line in value.split('\n'))
                height = max(height, lines * 16 + 12)
            if height > 409:
                raise ValueError(f'{title} row {r} is too long to display; split issue details into rows')
            ws.row_dimensions[r].height = height
            if title == 'Checklist' and r > 1:
                color = {'√': 'E7F2E9', 'X': 'FCE4E4', 'N/A': 'EEEEEE', 'Human check': 'FFF1CC'}[values[2]]
                ws.cell(r, 3).fill = PatternFill('solid', fgColor=color)
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:1'
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    payload = buffer.getvalue()
    def verify(source):
        check = load_workbook(source)
        try:
            if check.sheetnames != list(rows):
                raise ValueError('Workbook sheet mismatch')
            for title, expected in rows.items():
                actual = [list(row) for row in check[title].iter_rows(values_only=True)]
                if actual != expected or any(c.data_type == 'f' for row in check[title] for c in row):
                    raise ValueError('Workbook content verification failed')
        finally:
            check.close()
    verify(BytesIO(payload))
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open('xb') as out:
        out.write(payload)
    verify(target)
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output-dir', required=True, type=Path)
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding='utf-8-sig'))
        print(build_report(data, args.output_dir))
    except (ValueError, OSError, ImportError) as exc:
        parser.exit(1, f'Report failed: {exc}\n')


if __name__ == '__main__':
    main()
